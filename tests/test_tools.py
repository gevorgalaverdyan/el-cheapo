from datetime import date, datetime, timezone
from decimal import Decimal

import pytest

from elcheapo.agent.tools import make_tools
from elcheapo.models import Expense
from elcheapo.store.memory import InMemoryRepository

LOGGED = datetime(2026, 9, 12, tzinfo=timezone.utc)


def an_expense(day: int, amount: str, category: str, merchant: str = "", note: str = ""):
    return Expense(
        date=date(2026, 9, day),
        amount=Decimal(amount),
        category=category,
        merchant=merchant,
        note=note,
        source="text",
        logged_at=LOGGED,
        draft_id=f"d{day}{amount}",
    )


@pytest.fixture
def tools():
    repository = InMemoryRepository(categories=["Dining", "Groceries", "Transport"])
    for expense in [
        an_expense(1, "12.00", "Dining", "Tim Hortons", "morning coffee"),
        an_expense(3, "85.40", "Groceries", "Loblaws"),
        an_expense(5, "45.00", "Dining", "Kinton Ramen"),
        an_expense(8, "62.30", "Transport", "Petro-Canada"),
    ]:
        repository.append_expense(expense)
    repository.add_category("Vet")
    return {tool.__name__: tool for tool in make_tools(repository)}, repository


def test_the_agent_is_given_the_expected_tools(tools):
    by_name, _ = tools

    assert set(by_name) == {
        "propose_expense",
        "list_categories",
        "query_expenses",
        "export_expenses",
    }


def test_list_categories_returns_platform_and_user_categories(tools):
    by_name, _ = tools

    names = [c["name"] for c in by_name["list_categories"]()["categories"]]

    assert names == ["Dining", "Groceries", "Transport", "Vet"]


def test_list_categories_says_which_are_the_users_own(tools):
    by_name, _ = tools

    scopes = {c["name"]: c["scope"] for c in by_name["list_categories"]()["categories"]}

    assert scopes["Dining"] == "platform"
    assert scopes["Vet"] == "user"


def test_query_expenses_with_no_filters_returns_everything(tools):
    by_name, _ = tools

    result = by_name["query_expenses"]()

    assert result["count"] == 4


def test_query_expenses_reports_the_total(tools):
    by_name, _ = tools

    result = by_name["query_expenses"]()

    assert result["total"] == "204.70"


def test_query_expenses_totals_only_what_matched(tools):
    by_name, _ = tools

    result = by_name["query_expenses"](category="Dining")

    assert result["count"] == 2
    assert result["total"] == "57.00"


def test_query_expenses_filters_by_merchant(tools):
    by_name, _ = tools

    assert by_name["query_expenses"](merchant="tim")["count"] == 1


def test_query_expenses_filters_by_date_range(tools):
    by_name, _ = tools

    result = by_name["query_expenses"](date_from="2026-09-03", date_to="2026-09-05")

    assert result["count"] == 2


def test_query_expenses_returns_json_safe_amounts(tools):
    by_name, _ = tools

    first = by_name["query_expenses"]()["expenses"][0]

    # Decimal is not JSON-serialisable, and the agent's reply must serialise.
    assert isinstance(first["amount"], str)
    assert isinstance(first["date"], str)


def test_query_expenses_respects_a_limit(tools):
    by_name, _ = tools

    assert by_name["query_expenses"](limit=2)["count"] == 2


def test_a_malformed_date_is_reported_not_raised(tools):
    by_name, _ = tools

    result = by_name["query_expenses"](date_from="last tuesday")

    # The agent should get a message it can act on, not a stack trace that
    # ends the turn.
    assert "error" in result
    assert "date_from" in result["error"]


def test_a_malformed_amount_is_reported_not_raised(tools):
    by_name, _ = tools

    result = by_name["query_expenses"](min_amount="a lot")

    assert "error" in result


def test_a_query_matching_nothing_is_an_empty_result_not_an_error(tools):
    by_name, _ = tools

    result = by_name["query_expenses"](category="Rent")

    assert result["count"] == 0
    assert result["total"] == "0.00"
    assert result["expenses"] == []


# --- export ------------------------------------------------------------

@pytest.fixture
def exporting():
    repository = InMemoryRepository(categories=["Dining", "Groceries"])
    for expense in [
        an_expense(1, "12.00", "Dining", "Tim Hortons"),
        an_expense(3, "85.40", "Groceries", "Loblaws"),
        an_expense(5, "45.00", "Dining", "Kinton Ramen"),
    ]:
        repository.append_expense(expense)
    documents = []
    tools = {t.__name__: t for t in make_tools(repository, currency="CAD", documents=documents)}
    return tools["export_expenses"], documents


def test_export_is_offered_as_a_tool(tools):
    by_name, _ = tools

    assert "export_expenses" in by_name


def test_exporting_produces_a_spreadsheet_by_default(exporting):
    export, documents = exporting

    export()

    assert documents[0].filename.endswith(".xlsx")


def test_exporting_as_csv_when_asked(exporting):
    export, documents = exporting

    export(format="csv")

    assert documents[0].filename.endswith(".csv")
    assert documents[0].data.startswith(b"\xef\xbb\xbf")  # BOM for Excel


def test_the_export_reports_what_it_contains(exporting):
    export, _ = exporting

    result = export()

    assert result["rows"] == 3
    assert result["total"] == "142.40"


def test_filters_narrow_the_export(exporting):
    export, _ = exporting

    result = export(category="Dining")

    assert result["rows"] == 2
    assert result["total"] == "57.00"


def test_the_agent_can_name_the_file(exporting):
    export, documents = exporting

    export(filename="september dining")

    assert documents[0].filename == "september-dining.xlsx"


def test_a_dangerous_filename_is_made_safe(exporting):
    export, documents = exporting

    export(filename="../../etc/passwd")

    assert "/" not in documents[0].filename
    assert ".." not in documents[0].filename


def test_an_unknown_format_is_reported_not_raised(exporting):
    export, documents = exporting

    result = export(format="pdf")

    assert "error" in result
    assert documents == []


def test_an_unknown_grouping_is_reported(exporting):
    export, _ = exporting

    assert "error" in export(group_by="colour")


def test_exporting_nothing_explains_itself_instead_of_sending_an_empty_file(exporting):
    export, documents = exporting

    result = export(category="Rent")

    assert "error" in result
    assert documents == []


def test_grouping_is_passed_through_to_the_report(exporting):
    import io
    from openpyxl import load_workbook

    export, documents = exporting
    export(group_by="category")

    book = load_workbook(io.BytesIO(documents[0].data))
    assert "By category" in book.sheetnames
