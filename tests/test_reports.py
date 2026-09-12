import csv
import io
from datetime import date, datetime, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from elcheapo.models import Expense
from elcheapo.reports import build_csv, build_workbook, build_xlsx

LOGGED = datetime(2026, 9, 12, tzinfo=timezone.utc)


def an_expense(day: int, amount: str, category: str, merchant: str = "", note: str = ""):
    return Expense(
        date=date(2026, 9, day) if day <= 30 else date(2026, 8, day - 30),
        amount=Decimal(amount),
        category=category,
        merchant=merchant,
        note=note,
        source="text",
        logged_at=LOGGED,
        draft_id=f"d{day}{amount}",
    )


@pytest.fixture
def expenses() -> list[Expense]:
    return [
        an_expense(1, "12.00", "Dining", "Tim Hortons", "coffee"),
        an_expense(3, "85.40", "Groceries", "Loblaws"),
        an_expense(5, "45.00", "Dining", "Kinton Ramen"),
        an_expense(38, "62.30", "Transport", "Petro-Canada"),  # August
    ]


def rows_of(data: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))


# --- CSV ---------------------------------------------------------------

def test_csv_starts_with_a_header(expenses):
    assert rows_of(build_csv(expenses))[0] == [
        "date", "amount", "category", "merchant", "note", "source"
    ]


def test_csv_has_one_line_per_expense(expenses):
    assert len(rows_of(build_csv(expenses))) == 5  # header + 4


def test_csv_amounts_keep_two_decimal_places(expenses):
    amounts = [row[1] for row in rows_of(build_csv(expenses))[1:]]

    assert "12.00" in amounts


def test_csv_grouped_by_category_totals_each_group(expenses):
    rows = rows_of(build_csv(expenses, group_by="category"))
    totals = {row[0]: row[2] for row in rows[1:]}

    assert totals["Dining"] == "57.00"
    assert totals["Groceries"] == "85.40"


def test_csv_grouped_by_category_counts_each_group(expenses):
    rows = rows_of(build_csv(expenses, group_by="category"))
    counts = {row[0]: row[1] for row in rows[1:]}

    assert counts["Dining"] == "2"


def test_csv_grouped_by_month_splits_across_months(expenses):
    rows = rows_of(build_csv(expenses, group_by="month"))
    months = {row[0] for row in rows[1:]}

    assert months == {"2026-08", "2026-09"}


def test_csv_grouped_by_merchant(expenses):
    rows = rows_of(build_csv(expenses, group_by="merchant"))

    assert "Loblaws" in {row[0] for row in rows[1:]}


def test_csv_of_nothing_is_still_a_valid_file_with_headers():
    rows = rows_of(build_csv([]))

    assert len(rows) == 1


# --- XLSX --------------------------------------------------------------

def test_xlsx_always_has_an_expenses_sheet(expenses):
    book = load_workbook(io.BytesIO(build_xlsx(expenses, currency="CAD")))

    assert "Expenses" in book.sheetnames


def test_xlsx_detail_rows_match_the_expenses(expenses):
    book = load_workbook(io.BytesIO(build_xlsx(expenses, currency="CAD")))
    sheet = book["Expenses"]

    assert sheet.max_row == 5  # header + 4


def test_xlsx_amounts_are_numbers_not_text(expenses):
    """Text amounts cannot be summed in Excel, which defeats the point."""
    book = load_workbook(io.BytesIO(build_xlsx(expenses, currency="CAD")))

    assert isinstance(book["Expenses"].cell(row=2, column=2).value, (int, float))


def test_xlsx_grouping_adds_a_summary_sheet(expenses):
    book = load_workbook(
        io.BytesIO(build_xlsx(expenses, currency="CAD", group_by="category"))
    )

    assert "By category" in book.sheetnames


def test_xlsx_without_grouping_has_only_the_detail_sheet(expenses):
    book = load_workbook(io.BytesIO(build_xlsx(expenses, currency="CAD")))

    assert book.sheetnames == ["Expenses"]


def test_xlsx_summary_carries_a_chart(expenses):
    book = build_workbook(expenses, currency="CAD", group_by="category")

    assert book["By category"]._charts


def test_a_chart_can_be_suppressed(expenses):
    book = build_workbook(
        expenses, currency="CAD", group_by="category", include_chart=False
    )

    assert not book["By category"]._charts


def test_xlsx_of_nothing_still_opens(expenses):
    book = load_workbook(io.BytesIO(build_xlsx([], currency="CAD")))

    assert book["Expenses"].max_row == 1
